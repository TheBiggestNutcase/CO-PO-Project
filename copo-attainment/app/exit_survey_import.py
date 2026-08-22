"""
Best-effort extraction of Course Exit Survey responses from a Google
Sheets export of a Google Form's responses (.xlsx or .csv), so ratings
can be pre-filled instead of typed in one by one on the manual grid
(marks/exit_survey_grid.html).

Assumed shape (the typical Google Forms "Form responses 1" sheet): one
row per submission, one column per form question, plus a Timestamp
column and a column where each student typed their own USN. Like
roster_import.py and syllabus_import.py, this only produces a
best-effort *parse* - the caller always routes it through a human-
editable review screen (see the exit_survey_import_* routes in
app/routes/marks_routes.py) before anything is saved. Two things in
particular need a human's judgement, not just heuristics, and are only
ever offered here as an editable starting guess:

  - Matching a sheet column (a full Google Form question, e.g. "How
    would you rate the pace of the course?") to one of this course's
    already-set-up exit-survey AssessmentItems - those only carry a
    short label like "Q3", not the full question text, so there's no
    text to match against. The only guess on offer is positional
    (sheet column order == item seq order), which is usually right
    since a form is normally built in the same order as the item list,
    but it's still just a guess.
  - A rating option's exact wording is whatever the form's author
    typed, and different forms use genuinely different five-point
    scales - "Excellent/Good/Average/Poor/Very Poor" and "Excellent/
    Very Good/Good/Satisfactory/Not Satisfactory" have both shown up in
    real exit-survey forms for this app, and they disagree about what
    bare "Good" means (B in the first scale, C in the second, since
    "Very Good" already occupies B there). _map_rating below recognizes
    both, resolving the ambiguous word per-column by checking whether
    that column's other answers include a scale-specific marker word
    ("Very Good" or "Satisfactory"). Anything it still doesn't
    recognize is left unmapped rather than guessed - and because a
    whole column tends to share one form's wording, the review screen
    offers a per-distinct-answer bulk fix (see ParsedQuestionColumn.
    value_counts) instead of making someone fix the same word by hand
    on every row it appears in.
"""
import csv
import io
import re
from collections import Counter
from dataclasses import dataclass, field
from typing import Optional

from app.models import EXIT_SURVEY_RATINGS
from app.roster_import import USN_RE

_TIMESTAMP_HEADER_RE = re.compile(r"timestamp", re.IGNORECASE)
_USN_HEADER_RE = re.compile(r"\busn\b", re.IGNORECASE)

# Google Forms adds "Email Address" itself when "Collect email addresses"
# is on, and institutional forms commonly ask for Name/Section/Class/
# Branch/Roll Number alongside the USN - none of that is a rating
# question, but nothing marks it as such except that it *reads* like a
# short identifying-info label rather than a question. Guarded by length/
# "?" in _looks_like_metadata_header so this can't misfire on an actual
# question that happens to mention one of these words in passing.
_METADATA_HEADER_KEYWORDS = ("email", "e-mail", "name", "section", "class", "branch", "roll")


def _looks_like_metadata_header(header):
    h = header.strip().lower()
    if len(h) > 30 or "?" in h:
        return False  # long / question-shaped headers are never identifying-info fields
    return any(re.search(rf"\b{re.escape(kw)}\b", h) for kw in _METADATA_HEADER_KEYWORDS)


def _looks_like_free_text_column(raw_values):
    """A backstop for identifying-info columns a form phrases in a way
    _looks_like_metadata_header doesn't recognize: a real rating question
    only ever has a handful of distinct answers repeated across many
    rows, so a column where answers are almost all different from each
    other (like a name or email column) is very unlikely to be one."""
    values = [v.strip() for v in raw_values if (v or "").strip()]
    if len(values) < 5:
        return False  # too little data to tell anything from
    distinct = len({v.lower() for v in values})
    return distinct > 8 and distinct >= 0.5 * len(values)

_LEADING_LETTER_RE = re.compile(r"^\s*([A-E])\b")

# Words/phrases whose rank is the same in every known scale, checked in
# this order so a longer phrase is matched before a shorter one it
# contains ("not satisfactory" before "satisfactory", "very poor" before
# "poor"). Bare "good" is deliberately NOT here - see _resolve_good_letter.
_UNAMBIGUOUS_PHRASES = (
    ("not satisfactory", "E"),
    ("very poor", "E"),
    ("excellent", "A"),
    ("very good", "B"),
    ("satisfactory", "D"),
    ("average", "C"),
    ("poor", "D"),
)


def _resolve_good_letter(column_raw_values):
    """Bare "Good" is ambiguous between the two known five-point scales -
    decide which one a column is using by checking whether any of that
    column's own answers carry a marker unique to the expanded scale
    (Very Good / Satisfactory). Defaults to B, the simpler/older scale,
    when there's no such marker to go on."""
    for v in column_raw_values:
        lowered = (v or "").strip().lower()
        if "very good" in lowered or "satisfactory" in lowered:
            return "C"
    return "B"


def _map_rating(raw, good_letter="B"):
    """Best-effort free text -> A-E mapping. None means "not recognized" -
    the caller must flag that for a human to fix, never guess further.
    `good_letter` is the per-column resolution of bare "Good" from
    _resolve_good_letter."""
    if raw is None:
        return None
    text = raw.strip()
    if not text:
        return None
    upper = text.upper()
    if upper in EXIT_SURVEY_RATINGS:
        return upper
    m = _LEADING_LETTER_RE.match(upper)
    if m:
        return m.group(1)
    lowered = text.lower()
    for phrase, letter in _UNAMBIGUOUS_PHRASES:
        if phrase in lowered:
            return letter
    if "good" in lowered:
        return good_letter
    return None


@dataclass
class ParsedValueCount:
    raw: str  # a distinct answer text as it appears in this column
    letter: Optional[str]  # this column's current auto-mapped letter for that text
    count: int  # how many (surviving, post-dedup) rows have this exact text in this column


@dataclass
class ParsedQuestionColumn:
    col_index: int  # index into a raw sheet row (not position among question columns)
    header: str  # the sheet's own column header text (the Google Form question)
    suggested_item_id: Optional[int] = None  # positional guess, or None
    # Every distinct answer text seen in this column, most common first -
    # drives the review screen's "map this text to this letter, for every
    # row" bulk-fix control (a whole column tends to share one form's
    # wording, so fixing one word there fixes every row that used it).
    value_counts: list = field(default_factory=list)


@dataclass
class ParsedResponseCell:
    raw: str
    letter: Optional[str]  # mapped A-E, or None if unrecognized


@dataclass
class ParsedResponseRow:
    row_number: int  # 1-based among data rows (header excluded) - stable id for review-screen fields
    usn_raw: str
    matched_student_id: Optional[int] = None
    cells: dict = field(default_factory=dict)  # col_index -> ParsedResponseCell


@dataclass
class ParsedExitSurvey:
    usn_header: Optional[str] = None
    columns: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    warnings: list = field(default_factory=list)


def parse_exit_survey_file(file_obj, filename, existing_students, existing_items):
    """Parses an uploaded Google Sheets export of exit-survey responses
    into a ParsedExitSurvey.

    existing_students: the course's Student list (for USN matching).
    existing_items: the course's exit-survey AssessmentItem list, in seq
    order (for the positional column-matching guess).

    Raises ValueError for a file type this doesn't handle (caller should
    catch and show a friendly message); any other read failure (a
    corrupt file) propagates as-is, same contract as
    roster_import.parse_roster_docx and syllabus_import.parse_syllabus_pdf.
    """
    name = (filename or "").lower()
    if name.endswith(".csv"):
        rows = _read_csv(file_obj)
    elif name.endswith(".xlsx"):
        rows = _read_xlsx(file_obj)
    else:
        raise ValueError("Unsupported file type - upload the sheet as .xlsx or .csv.")

    result = ParsedExitSurvey()

    if not rows:
        result.warnings.append("The file appears to be empty.")
        return result

    header = [(c or "").strip() for c in rows[0]]
    data_rows = rows[1:]

    usn_col_idx = _find_usn_column(header, data_rows)
    if usn_col_idx is None:
        result.warnings.append(
            "Couldn't find a column that looks like it holds USNs - match each response to a "
            "student by hand on the review screen below."
        )
    else:
        result.usn_header = header[usn_col_idx]

    question_col_indices = []
    excluded_as_identifying_info = []
    for i, h in enumerate(header):
        if i == usn_col_idx or not h or _TIMESTAMP_HEADER_RE.search(h):
            continue
        if _looks_like_metadata_header(h):
            excluded_as_identifying_info.append(h)
            continue
        col_values = [row[i] for row in data_rows if i < len(row)]
        if _looks_like_free_text_column(col_values):
            excluded_as_identifying_info.append(h)
            continue
        question_col_indices.append(i)

    if excluded_as_identifying_info:
        result.warnings.append(
            "Treated these column(s) as identifying info (email/name/section-type fields), not "
            "rating questions, and left them out: " + ", ".join(excluded_as_identifying_info) + ". "
            "If one of these is actually a survey question, this heuristic guessed wrong - let me know."
        )

    for pos, col_idx in enumerate(question_col_indices):
        suggested_item_id = existing_items[pos].id if pos < len(existing_items) else None
        result.columns.append(ParsedQuestionColumn(
            col_index=col_idx, header=header[col_idx], suggested_item_id=suggested_item_id,
        ))

    students_by_usn = {s.usn.strip().upper(): s for s in existing_students}

    # Resolve each column's meaning of bare "Good" up front, from that
    # column's own full set of answers - needs a look at every row before
    # any row can be mapped, so this is a separate pass over data_rows.
    good_letter_by_col = {
        col.col_index: _resolve_good_letter(
            row[col.col_index] for row in data_rows if col.col_index < len(row)
        )
        for col in result.columns
    }

    parsed_rows = []
    last_row_number_by_usn = {}  # usn (upper) -> row_number of its last occurrence in the sheet

    for i, raw_row in enumerate(data_rows, start=1):
        if not any((c or "").strip() for c in raw_row):
            continue  # a blank trailing row - common in sheet exports

        usn_raw = raw_row[usn_col_idx].strip() if usn_col_idx is not None and usn_col_idx < len(raw_row) else ""
        usn_key = usn_raw.upper()
        matched_student = students_by_usn.get(usn_key) if usn_key else None

        prow = ParsedResponseRow(
            row_number=i, usn_raw=usn_raw,
            matched_student_id=matched_student.id if matched_student else None,
        )
        for col in result.columns:
            raw_val = raw_row[col.col_index] if col.col_index < len(raw_row) else ""
            letter = _map_rating(raw_val, good_letter=good_letter_by_col[col.col_index])
            prow.cells[col.col_index] = ParsedResponseCell(raw=raw_val, letter=letter)

        if usn_key:
            last_row_number_by_usn[usn_key] = i

        parsed_rows.append(prow)

    duplicate_usns = {
        usn for usn, count in
        _count_by(r.usn_raw.strip().upper() for r in parsed_rows if r.usn_raw.strip()).items()
        if count > 1
    }
    if duplicate_usns:
        result.warnings.append(
            f"{len(duplicate_usns)} USN(s) appear more than once in the sheet "
            f"({', '.join(sorted(duplicate_usns)[:10])}{', ...' if len(duplicate_usns) > 10 else ''}) - "
            "only the most recent submission for each is kept below."
        )
        parsed_rows = [
            r for r in parsed_rows
            if not r.usn_raw.strip() or r.row_number == last_row_number_by_usn[r.usn_raw.strip().upper()]
        ]

    result.rows = parsed_rows

    for col in result.columns:
        counts = Counter()
        letter_for_raw = {}
        for r in parsed_rows:
            cell = r.cells.get(col.col_index)
            if cell and cell.raw.strip():
                counts[cell.raw] += 1
                letter_for_raw[cell.raw] = cell.letter
        col.value_counts = sorted(
            (ParsedValueCount(raw=raw, letter=letter_for_raw[raw], count=n) for raw, n in counts.items()),
            key=lambda vc: -vc.count,
        )

    unmatched = sum(1 for r in parsed_rows if r.matched_student_id is None)
    if unmatched:
        result.warnings.append(
            f"{unmatched} response(s) couldn't be matched to a roster USN automatically - "
            "match them by hand below, or leave them unchecked to skip."
        )

    unmapped_cells = sum(
        1 for r in parsed_rows for c in r.cells.values() if c.raw.strip() and c.letter is None
    )
    if unmapped_cells:
        result.warnings.append(
            f"{unmapped_cells} answer(s) didn't match a recognized rating (recognizes Excellent/Good/"
            "Average/Poor/Very Poor and Excellent/Very Good/Good/Satisfactory/Not Satisfactory, or bare "
            "A-E) - use each column's \"map this answer\" list below to fix them all at once."
        )

    return result


def _count_by(values):
    counts = {}
    for v in values:
        counts[v] = counts.get(v, 0) + 1
    return counts


def _read_csv(file_obj):
    raw = file_obj.read()
    text = raw.decode("utf-8-sig") if isinstance(raw, bytes) else raw  # Google's CSV export includes a BOM
    return list(csv.reader(io.StringIO(text)))


def _read_xlsx(file_obj):
    from openpyxl import load_workbook

    wb = load_workbook(file_obj, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    return [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]


def _find_usn_column(header, data_rows):
    for i, h in enumerate(header):
        if _USN_HEADER_RE.search(h):
            return i
    # No header explicitly says "USN" - fall back to whichever column's
    # values look most like actual USNs.
    best_idx, best_score = None, 0
    for i in range(len(header)):
        sample = [row[i].strip() for row in data_rows[:50] if i < len(row) and row[i]]
        if not sample:
            continue
        score = sum(1 for v in sample if USN_RE.match(v.upper()))
        # A majority of the sampled cells in this column need to look like
        # a real USN - no fixed minimum count, so this also works against
        # the tiny sheets a test (or a brand-new course) might have.
        if score > best_score and score * 2 >= len(sample):
            best_idx, best_score = i, score
    return best_idx
