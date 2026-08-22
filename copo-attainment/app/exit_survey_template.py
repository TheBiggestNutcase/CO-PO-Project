"""
Builds a downloadable "exit survey question template" workbook: a ready-
made list of suggested exit-survey questions (matched to a course's own
Course Outcomes, if any are defined) plus instructions for turning them
into a Google Form, for a coordinator who doesn't already have one set up.

This is a static, formula-free workbook (no calculations to get right,
unlike export_xlsx.py's report) - it's just structured suggested content,
generated fresh each time from the course's current COs so the questions
stay in sync if COs change.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

RATING_SCALE = (
    ("A", "Excellent"),
    ("B", "Good"),
    ("C", "Average"),
    ("D", "Poor"),
    ("E", "Very Poor"),
)

GENERIC_QUESTIONS = (
    "How would you rate the pace at which this course was taught?",
    "How would you rate the clarity of the faculty's explanations?",
    "How would you rate the usefulness of the assignments/self-study work?",
    "How would you rate the fairness of the internal assessments?",
    "How would you rate the learning resources provided (notes, references, etc.)?",
    "Overall, how would you rate this course?",
)


def build_template_workbook(course):
    """Returns an openpyxl Workbook: an Instructions sheet and a Suggested
    Questions sheet, one row per course CO (worded around that CO) if the
    course has any, else a handful of generic questions."""
    wb = Workbook()

    instructions = wb.active
    instructions.title = "Instructions"
    _fill_instructions(instructions, course)

    questions = wb.create_sheet("Suggested Questions")
    _fill_questions(questions, course)

    return wb


def _fill_instructions(ws, course):
    bold = Font(name="Arial", bold=True, size=12)
    normal = Font(name="Arial", size=11)
    wrap = Alignment(wrap_text=True, vertical="top")

    lines = [
        (f"Exit survey template - {course.subject_code or 'this course'}", bold),
        ("", normal),
        ("1. Create a new Google Form (forms.google.com).", normal),
        ("2. Add a short-answer question titled exactly \"USN\" and mark it required - "
         "this is what lets the app match each response back to a student on the roster.", normal),
        ("3. For each row on the \"Suggested Questions\" sheet, add one \"Multiple choice\" "
         "question using that row's text, with exactly these 5 choices:", normal),
    ]
    for letter, word in RATING_SCALE:
        lines.append((f"      {letter} - {word}", normal))
    lines += [
        ("", normal),
        ("   Edit the suggested wording however you like - the exact text doesn't matter to the "
         "app, only the A-E choice scale does.", normal),
        ("4. Share the form with students and collect responses.", normal),
        ("5. In the form's \"Responses\" tab, click the Sheets icon to create a linked spreadsheet, "
         "then File > Download > Microsoft Excel (.xlsx) or File > Download > Comma Separated Values (.csv).", normal),
        ("6. Back in this app, open this course's Exit Survey marks page and use \"Import from file\" "
         "to upload that download - it will walk you through matching columns to questions and rows to "
         "students before saving anything.", normal),
    ]

    ws.column_dimensions["A"].width = 100
    for i, (text, font) in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.font = font
        cell.alignment = wrap


def _fill_questions(ws, course):
    bold = Font(name="Arial", bold=True, size=11)
    header = ["CO", "Suggested question", "Choices (use exactly these 5)"]
    for col, text in enumerate(header, start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font = bold

    choices_text = "; ".join(f"{letter} - {word}" for letter, word in RATING_SCALE)

    cos = sorted(course.outcomes, key=lambda co: co.seq) if course.outcomes else []
    row = 2
    if cos:
        for co in cos:
            desc = co.description or co.code
            question = f"How would you rate this course in helping you achieve {co.code} ({desc})?"
            ws.cell(row=row, column=1, value=co.code)
            ws.cell(row=row, column=2, value=question)
            ws.cell(row=row, column=3, value=choices_text)
            row += 1
    else:
        for question in GENERIC_QUESTIONS:
            ws.cell(row=row, column=1, value="")
            ws.cell(row=row, column=2, value=question)
            ws.cell(row=row, column=3, value=choices_text)
            row += 1

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 45
    for r in range(2, row):
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")


def build_template_bytes(course):
    wb = build_template_workbook(course)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
